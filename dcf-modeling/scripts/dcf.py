#!/usr/bin/env python3
"""Explicit end-year UFCF DCF; all amounts must use the same currency/unit."""

import argparse
import json
import math


def finite(value, name):
    if (
        isinstance(value, bool)
        or not isinstance(value, (int, float))
        or not math.isfinite(value)
    ):
        raise ValueError(f"{name} must be finite")
    return value


def calculate(model):
    if not model.get("currency") or not model.get("unit"):
        raise ValueError("currency and unit required; use shares in corresponding unit")
    shares = finite(model["diluted_shares"], "diluted_shares")
    if shares <= 0:
        raise ValueError("diluted_shares must be positive")
    net_debt = finite(model["net_debt"], "net_debt")
    price = finite(model["price"], "price") if model.get("price") is not None else None
    if price is not None and price <= 0:
        raise ValueError("price must be positive")
    results = {}
    if set(model["scenarios"]) != {"base", "bull", "bear"}:
        raise ValueError("base, bull, bear scenarios required")
    for name, scenario in model["scenarios"].items():
        flows = scenario["ufcf"]
        wacc = finite(scenario["wacc"], "wacc")
        g = finite(scenario["terminal_growth"], "terminal_growth")
        if not flows or not isinstance(flows, list):
            raise ValueError("Annual UFCF list required")
        if not -1 < g < wacc or not 0 < wacc < 1:
            raise ValueError("Require 0 < WACC < 1 and -1 < g < WACC")
        flows = [finite(v, "ufcf") for v in flows]
        if flows[-1] <= 0:
            raise ValueError("Perpetuity model requires positive terminal UFCF")
        pv = sum(f / (1 + wacc) ** (i + 1) for i, f in enumerate(flows))
        tv = flows[-1] * (1 + g) / (wacc - g)
        pv_tv = tv / (1 + wacc) ** len(flows)
        ev = pv + pv_tv
        equity = ev - net_debt
        per_share = equity / shares
        results[name] = {
            "enterprise_value": ev,
            "equity_value": equity,
            "value_per_share": per_share,
            "upside": per_share / price - 1 if price else None,
            "terminal_value_share": pv_tv / ev if ev else None,
            "assumptions": scenario,
        }
    return {
        "currency": model["currency"],
        "unit": model["unit"],
        "discount_convention": "end-of-year, annual UFCF, no stub period",
        "is_demo": model.get("is_demo", False),
        "results": results,
    }


def workbook(model, output):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "DCF"
    ws.append(["Assumption", "Base", "Bull", "Bear"])
    ws.append(
        ["WACC"] + [model["scenarios"][s]["wacc"] for s in ("base", "bull", "bear")]
    )
    ws.append(
        ["Terminal growth"]
        + [model["scenarios"][s]["terminal_growth"] for s in ("base", "bull", "bear")]
    )
    ws.append(["Net debt"] + [model["net_debt"]] * 3)
    ws.append(["Diluted shares"] + [model["diluted_shares"]] * 3)
    n = max(len(s["ufcf"]) for s in model["scenarios"].values())
    for i in range(n):
        ws.append(
            [f"Year {i + 1} UFCF"]
            + [
                model["scenarios"][s]["ufcf"][i]
                if i < len(model["scenarios"][s]["ufcf"])
                else None
                for s in ("base", "bull", "bear")
            ]
        )
    row = 6 + n
    ws.cell(row, 1, "Value per share (formula)")
    for col, s in zip("BCD", ("base", "bull", "bear")):
        horizon = len(model["scenarios"][s]["ufcf"])
        pv = "+".join(f"{col}{6 + i}/(1+{col}2)^{i + 1}" for i in range(horizon))
        ws[f"{col}{row}"] = (
            f"=(({pv})+{col}{5 + horizon}*(1+{col}3)/({col}2-{col}3)/(1+{col}2)^{horizon}-{col}4)/{col}5"
        )
    ws.append(
        ["Python-verified values"]
        + [
            calculate(model)["results"][s]["value_per_share"]
            for s in ("base", "bull", "bear")
        ]
    )
    ws.append(["Currency / amount unit", model["currency"], model["unit"]])
    ws.append(
        ["Formula caches not recalculated by openpyxl; open in Excel or LibreOffice."]
    )
    ws.append(
        [
            "Fictional example"
            if model.get("is_demo")
            else "User-supplied assumptions; not an investment recommendation"
        ]
    )
    ws.column_dimensions["A"].width = 65
    for row_cells in ws.iter_rows(min_row=2, max_row=5 + n, min_col=2, max_col=4):
        for cell in row_cells:
            cell.font = Font(color="0000FF")
    wb.save(output)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("input")
    parser.add_argument("--xlsx")
    args = parser.parse_args()
    with open(args.input) as f:
        model = json.load(f)
    result = calculate(model)
    if args.xlsx:
        workbook(model, args.xlsx)
        result["workbook_note"] = (
            "Excel formulas require recalculation when opened; Python values verified independently."
        )
    print(json.dumps(result, ensure_ascii=False, indent=2, allow_nan=False))


if __name__ == "__main__":
    main()
